from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import csv, json, zipfile, shutil, os

OUT=Path('/mnt/data/way_v12')
SRC=OUT/'source'

TITLE='THE WAY'
ORG='UNDECIDED'
DATE='September 21, 2026'

NAVY='1F2937'; SLATE='334155'; PALE='E2E8F0'; LIGHT='F8FAFC'; GOLD='D9B44A'; RED='B91C1C'; GREEN='166534'


def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=tcPr.find(qn('w:shd'))
    if shd is None:
        shd=OxmlElement('w:shd'); tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_text_color(cell, color, bold=False):
    for p in cell.paragraphs:
        for r in p.runs:
            r.font.color.rgb=RGBColor.from_string(color); r.bold=bold

def repeat_table_header(row):
    trPr=row._tr.get_or_add_trPr(); tblHeader=OxmlElement('w:tblHeader'); tblHeader.set(qn('w:val'),'true'); trPr.append(tblHeader)

def cant_split(row):
    trPr=row._tr.get_or_add_trPr(); x=OxmlElement('w:cantSplit'); trPr.append(x)

def keep_with_next(p):
    p.paragraph_format.keep_with_next=True

def set_repeat_and_nosplit(table):
    if table.rows:
        repeat_table_header(table.rows[0])
    for r in table.rows:
        cant_split(r)

def add_page_number(paragraph):
    paragraph.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    run=paragraph.add_run('Page ')
    fldChar1=OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'),'begin')
    instrText=OxmlElement('w:instrText'); instrText.set(qn('xml:space'),'preserve'); instrText.text=' PAGE '
    fldChar2=OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'),'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

def base_doc(doc_title, subtitle, version='1.2'):
    d=Document()
    sec=d.sections[0]
    sec.top_margin=Inches(.55); sec.bottom_margin=Inches(.55); sec.left_margin=Inches(.65); sec.right_margin=Inches(.65)
    styles=d.styles
    styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.4)
    styles['Normal'].paragraph_format.space_after=Pt(4)
    for style_name,size,color in [('Title',24,NAVY),('Heading 1',16,NAVY),('Heading 2',12,SLATE),('Heading 3',10.5,SLATE)]:
        st=styles[style_name]; st.font.name='Aptos Display' if style_name=='Title' else 'Aptos'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.font.bold=True
        st.paragraph_format.space_before=Pt(7); st.paragraph_format.space_after=Pt(4)
    # Header/footer
    h=sec.header.paragraphs[0]; h.text=f'{ORG} | The Way | {doc_title} v{version}'; h.style=styles['Normal']; h.runs[0].font.size=Pt(8); h.runs[0].font.color.rgb=RGBColor.from_string(SLATE)
    f=sec.footer.paragraphs[0]; f.text=f'{DATE} | '; f.runs[0].font.size=Pt(8); f.runs[0].font.color.rgb=RGBColor.from_string(SLATE); add_page_number(f)
    # Cover
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(TITLE); r.bold=True; r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string(GOLD)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(doc_title); r.bold=True; r.font.size=Pt(24); r.font.color.rgb=RGBColor.from_string(NAVY)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(subtitle); r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(SLATE)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(f'Version {version} | {DATE}'); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string(SLATE)
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run('Working legal/operational package for UNDECIDED - not a substitute for attorney or tax-professional review'); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string(RED)
    d.add_page_break()
    return d

def h1(d,text):
    p=d.add_paragraph(text,style='Heading 1'); keep_with_next(p); return p

def h2(d,text):
    p=d.add_paragraph(text,style='Heading 2'); keep_with_next(p); return p

def h3(d,text):
    p=d.add_paragraph(text,style='Heading 3'); keep_with_next(p); return p

def bullet(d,text,level=0):
    p=d.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p

def num(d,text):
    # manual numbering handled by caller for reliable restarts
    p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(.18); p.paragraph_format.first_line_indent=Inches(-.18); p.add_run(text); return p

def callout(d, label, text, fill='FFF7D6'):
    t=d.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=True
    c=t.cell(0,0); set_cell_shading(c,fill); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p=c.paragraphs[0]; r=p.add_run(label+': '); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); p.add_run(text)
    set_repeat_and_nosplit(t); d.add_paragraph()

def table(d, headers, rows, widths=None, font=8.4):
    t=d.add_table(rows=1, cols=len(headers)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.text=str(h); set_cell_shading(c,SLATE); set_cell_text_color(c,'FFFFFF',True); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for row in rows:
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            for p in cells[i].paragraphs:
                for r in p.runs: r.font.size=Pt(font)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths):
                row.cells[i].width=Inches(w)
    set_repeat_and_nosplit(t)
    d.add_paragraph()
    return t

def save_doc(d, filename):
    path=OUT/filename; d.save(path); return path

# ------------------ Doc 1: Incorporation and Board Execution ------------------
d=base_doc('Filing-Ready Incorporation & Founding Board Execution Packet','Tennessee charter data, corporate bylaws interface, incorporator action, initial board organization, officer appointments and signature instruments')

h1(d,'1. Status and execution boundary')
d.add_paragraph('This packet is designed to be substantially complete before the final legal name, additional directors, Secretary, and street addresses are known. Bracketed fields are controlled execution gates rather than missing drafting work.')
callout(d,'Current hard gates','Do not file until a clearable legal name is selected, at least three natural-person directors are identified for the corporation, a President and a different person as Secretary are available, and registered/principal office information is ready for the filing. Tennessee currently requires a board of at least three natural persons and requires President and Secretary to be different individuals.')

table(d,['Item','Current decision','Execution status'],[
['Entity type','Tennessee public benefit religious nonprofit corporation','WORKING DECISION'],
['Statutory members','No statutory members at formation','WORKING DECISION - confirm before filing'],
['Fiscal year','Calendar year ending December 31','FIXED'],
['Founder','Thomas Carter - Founder; intended President/Executive Director','PROVISIONAL UNTIL BOARD APPOINTMENT'],
['Board size','Legal floor 3; formation target 5; mature target 7-11','DESIGN BASELINE'],
['Board independence','Formation majority independent; mature target >= 2/3 independent','DESIGN BASELINE'],
['Legal name','[[LEGAL_NAME]], Inc.','OPEN'],
['Registered agent','Founder intended','READY SUBJECT TO ADDRESS'],
['Registered office','[[REGISTERED_OFFICE]]','OPEN'],
['Principal office','[[PRINCIPAL_OFFICE]]','OPEN'],
])

h1(d,'2. Tennessee charter filing data sheet')
d.add_paragraph('Use this sheet to populate the Tennessee Secretary of State\'s then-current domestic nonprofit charter filing interface/form. The state currently lists the domestic nonprofit charter (SS-4418) at a $100 filing fee; verify immediately before submission.')

table(d,['Field','Proposed entry / instruction'],[
['Corporate name','[[LEGAL_NAME]], Inc. - must pass Tennessee availability/confusion review.'],
['Type','Public benefit corporation; religious corporation where the form or attachment permits classification.'],
['Members','The corporation WILL NOT have members within the meaning of the Tennessee Nonprofit Corporation Act.'],
['Principal office','[[PRINCIPAL_OFFICE]]'],
['Registered agent','Thomas Carter'],
['Registered office','[[REGISTERED_OFFICE_TENNESSEE]]'],
['Incorporator','Thomas Carter - address [[INCORPORATOR_ADDRESS]]'],
['Initial directors','Either list in charter if chosen or elect through incorporator organizational action after filing. Final approach to be selected immediately before filing.'],
['Fiscal year','December 31'],
['Effective date','Upon filing unless a delayed effective date is specifically desired.'],
])

h2(d,'2.1 Proposed charter attachment language')
d.add_paragraph('The following is model substantive language for attorney review and attachment/inclusion where appropriate. It should be reconciled to the Secretary of State filing fields rather than duplicating fields unnecessarily.')

sections=[
('ARTICLE 1 - NAME','The name of the corporation is [[LEGAL_NAME]], Inc. (the "Corporation").'),
('ARTICLE 2 - CHARACTER','The Corporation is organized as a Tennessee nonprofit public benefit religious corporation. It shall have no members within the meaning of the Tennessee Nonprofit Corporation Act.'),
('ARTICLE 3 - PURPOSES','The Corporation is organized and shall be operated exclusively for religious, charitable, and educational purposes within the meaning of section 501(c)(3) of the Internal Revenue Code, including advancement of religion; public worship and prayer; religious and ethical formation; historical, linguistic, philosophical, and comparative study of religious traditions; publication and education; charitable assistance; community service; pastoral and spiritual care; development and support of Houses of Prayer and related religious communities; and any other lawful activities that further such exempt purposes.'),
('ARTICLE 4 - SECTION 501(c)(3) LIMITATIONS','No part of the net earnings of the Corporation shall inure to the benefit of, or be distributable to, any director, officer, founder, employee, or other private person, except that the Corporation is authorized to pay reasonable compensation for services actually rendered and to make payments and distributions in furtherance of its exempt purposes. No substantial part of the activities of the Corporation shall consist of carrying on propaganda or otherwise attempting to influence legislation except to the extent permitted for an organization described in section 501(c)(3). The Corporation shall not participate or intervene in any political campaign on behalf of or in opposition to any candidate for public office. The Corporation shall not conduct activities prohibited to an organization exempt under section 501(c)(3) or contributions to which are deductible under section 170(c)(2).'),
('ARTICLE 5 - FOUNDATIONAL RIGHTS AND ANTI-COERCION LIMITS','The Corporation shall operate in a manner consistent with freedom of conscience, voluntary participation, safe exit, safeguarding of children and vulnerable persons, non-retaliation for good-faith complaints or criticism, financial integrity, and the principle that no living founder, officer, minister, director, council, or other person possesses inherent infallibility or unlimited institutional authority. Claimed revelation or spiritual discernment shall not by itself override civil law, safeguarding duties, fiduciary controls, governing documents, or another adult\'s legal or personal consent.'),
('ARTICLE 6 - BOARD','The affairs of the Corporation shall be managed under the direction of a Board of Directors in accordance with the bylaws. The Board shall consist of not fewer than three natural persons. The bylaws may establish a larger fixed number or a range. The Corporation intends to maintain an independent governing structure and to limit concentration of financial, executive, and religious authority.'),
('ARTICLE 7 - FOUNDER STATUS','Founder status is historical only. It does not confer ownership, a property interest, permanent office, doctrinal infallibility, unilateral amendment authority, veto power, succession authority, entitlement to compensation, immunity from discipline, or a right to appoint a permanent board majority.'),
('ARTICLE 8 - DISSOLUTION','Upon dissolution, after paying or making provision for liabilities, the Corporation shall distribute its remaining assets exclusively for one or more purposes described in section 501(c)(3), to one or more organizations then recognized or qualifying for such purposes, or to a federal, state, or local government for a public purpose. No remaining assets shall be distributed to a founder, director, officer, member of the religious community, or other private person except in satisfaction of lawful obligations.'),
]
for title,text in sections:
    h3(d,title); d.add_paragraph(text)

h1(d,'3. Bylaws interface - provisions to lock before filing/organization')
table(d,['Bylaw element','v1.2 working rule'],[
['Board size','5 during formation where feasible; never fewer than 3; expansion to 7-11 within the founding period.'],
['Director terms','Three-year terms, staggered as practicable; maximum two consecutive full terms before a one-year break, subject to narrow emergency exception.'],
['Independence','At least a majority independent during formation; target >= 2/3 independent at mature structure.'],
['Selection without statutory members','Initial directors elected by incorporator or named in charter; later directors selected as the bylaws provide, with independent nominating process and board election.'],
['Officers','President, Secretary, Treasurer and such other officers as the Board creates. President and Secretary must be different persons.'],
['Chair','Independent Chair preferred; President/Executive Director should not serve as Chair after formation transition.'],
['Reserved independent powers','Executive compensation, related-party transactions, senior-leader investigations, safeguarding, external financial review, and entrenched-right amendments.'],
['Written action','Permit written board action only in a manner consistent with applicable Tennessee law and the bylaws; maintain signed consents with minutes.'],
['Companions vs members','Companion is a religious identity only and creates no statutory corporate membership or voting rights.'],
])

h1(d,'4. Incorporator organizational action - model')
d.add_paragraph('Use after the Tennessee charter has been accepted if the initial directors were not already fully established by the charter. Attorney should confirm the exact sequence used in the final filing.')
model=[
('Formation acknowledgement','The incorporator acknowledges filing and acceptance of the Charter of [[LEGAL_NAME]], Inc. by the Tennessee Secretary of State on [[FORMATION_DATE]].'),
('Election of initial directors','The incorporator elects the following initial directors, each having consented to serve: [[DIRECTOR_1]], [[DIRECTOR_2]], [[DIRECTOR_3]], [[DIRECTOR_4]], [[DIRECTOR_5]]. At least three must be natural persons; unused placeholders may be deleted if the initial lawful board consists of three or four persons.'),
('Delivery of records','The incorporator delivers the filed Charter, this action, and formation records to the initial Board and thereafter has no governance authority by virtue of incorporator status alone.'),
]
for i,(a,b) in enumerate(model,1): num(d,f'{i}. {a}. {b}')
d.add_paragraph('Incorporator: Thomas Carter   Signature: ______________________________   Date: __________________')

h1(d,'5. Initial Board organizational meeting / unanimous written consent')
d.add_paragraph('The following resolutions may be used at the first Board meeting or converted to a signed unanimous written consent when lawful and appropriate. Tennessee currently allows board action without a meeting unless the charter/bylaws provide otherwise, when all directors consent to acting without a meeting and the required affirmative vote is obtained, with written consents retained in the minutes.')

resolutions=[
'Ratify the filed Charter and formation acts.',
'Adopt the Bylaws and the selected nonmember governance schedule.',
'Ratify the Constitution of the Way v1.0 as the current religious constitutional baseline, expressly subordinate to civil law and the corporate Charter/Bylaws for civil governance.',
'Adopt the Teaching & Practice Baseline v1.0 and Companion Handbook v1.0 as current religious-operational documents subject to their review procedures.',
'Adopt the Institutional Integrity & Anti-Coercion Framework and associated conflicts, compensation, finance, safeguarding, complaints, investigations, discipline/appeals, exit/non-retaliation, pastoral ethics, revelation/doctrine, and external-review policies.',
'Appoint officers: President [[NAME]], Secretary [[NAME - MUST DIFFER FROM PRESIDENT]], Treasurer [[NAME]], and Chair [[NAME]].',
'Authorize application for an EIN after legal formation and designate the principal officer/responsible party for that application.',
'Authorize opening bank and payment accounts after EIN issuance under the approved authority matrix.',
'Approve a calendar fiscal year ending December 31.',
'Approve the conflict/independence disclosures of all directors and officers and record any recusals.',
'Establish the Governance & Nominating, Finance/Audit/Compensation, and Safeguarding & Institutional Integrity committees as staffing permits.',
'Authorize preparation and electronic submission of a full Form 1023 after final Board review of the activity narrative, financial projections, compensation arrangements, governing documents, and public-charity classification request.',
'Authorize a Founding Gathering Circle / House-in-Formation pilot subject to safeguarding, complaint-channel, finance, and public-claims controls.',
'Adopt a corporate records book and retention/security structure.',
]
for i,x in enumerate(resolutions,1): num(d,f'{i}. RESOLVED, that {x}')

h2(d,'5.1 Signature / vote record')
table(d,['Director','Independent?','Present / consent','For','Against','Abstain / recused','Signature'],[
['[[DIRECTOR_1]]','[[Y/N]]','','','','',''],['[[DIRECTOR_2]]','[[Y/N]]','','','','',''],['[[DIRECTOR_3]]','[[Y/N]]','','','','',''],['[[DIRECTOR_4]]','[[Y/N]]','','','','',''],['[[DIRECTOR_5]]','[[Y/N]]','','','','',''],
],font=7.5)

h1(d,'6. Officer acceptance instruments')
for role,limits in [
('President / Executive Director','Subject to Board supervision; no authority to self-approve compensation, related-party transactions, senior-leader investigations, or entrenched-right changes.'),
('Secretary','Responsible for minutes and authentication of corporate records as assigned; must be a different individual from the President.'),
('Treasurer','Oversees financial reporting and controls as delegated; does not displace Board fiduciary responsibility.'),
('Board Chair','Facilitates Board governance and executive oversight; preferred to be independent from executive staff.'),
]:
    h2(d,role)
    d.add_paragraph(f'I, [[NAME]], accept appointment as {role} of [[LEGAL_NAME]], Inc., subject to the Charter, Bylaws, lawful Board resolutions, fiduciary duties, conflict-of-interest requirements, safeguarding obligations, and the following role limit: {limits}')
    d.add_paragraph('Signature: ______________________________   Date: __________________')

h1(d,'7. Controlled filing gates')
table(d,['Gate','Required evidence','Status now'],[
['G1 - legal name','Tennessee availability/confusion review; web/trademark/domain review appropriate to intended public use.','OPEN'],
['G2 - directors','At least three natural persons willing and legally able to serve; formation target five.','OPEN - founder only'],
['G3 - Secretary','A person other than the President willing to serve.','OPEN'],
['G4 - addresses','Registered office in Tennessee; principal/mailing address; incorporator address.','PLACEHOLDERS'],
['G5 - charter counsel review','Review final purpose, member, rights, founder-limit, and dissolution language.','PENDING'],
['G6 - filing fee/interface','Verify current Tennessee fee and filing mechanism immediately before filing.','RECHECK AT FILING'],
['G7 - organizational meeting','Signed incorporator/director actions and conflict disclosures.','POST-FILING'],
['G8 - EIN','Apply only after legal formation; principal officer/responsible party identified.','POST-FILING'],
['G9 - bank/accounting','EIN, filed charter, Board authorization, signer matrix, bookkeeping.','POST-EIN'],
['G10 - Form 1023','Complete activity narrative, 3-year financials, compensation disclosures, governing documents, public-charity request, Pay.gov submission.','POST-BOARD'],
])

h1(d,'8. Current source notes')
sources=[
('Tennessee Secretary of State - Business Forms & Fees','Domestic nonprofit Charter SS-4418 currently listed at $100; name reservation currently $20.','https://sos.tn.gov/businesses/forms-and-fees'),
('Tenn. Code § 48-58-103 (2025)','Nonprofit board must consist of three or more natural persons.','https://law.justia.com/codes/tennessee/title-48/nonprofit-corporations/chapter-58/part-1/section-48-58-103/'),
('Tenn. Code § 48-58-401 (2025)','Every corporation must have President and Secretary; same person cannot hold both offices simultaneously.','https://law.justia.com/codes/tennessee/title-48/nonprofit-corporations/chapter-58/part-4/section-48-58-401/'),
('Tenn. Code § 48-58-104 (2025)','For a corporation without members, director selection follows charter/bylaws; absent another method, directors are elected by the Board.','https://law.justia.com/codes/tennessee/title-48/nonprofit-corporations/chapter-58/part-1/section-48-58-104/'),
('Tenn. Code § 48-58-202 (2025)','Board action without meeting mechanics.','https://law.justia.com/codes/tennessee/title-48/nonprofit-corporations/chapter-58/part-2/section-48-58-202/'),
('IRS Form 1023 Instructions','Organizing-document, public-charity and church-classification rules.','https://www.irs.gov/instructions/i1023'),
]
table(d,['Source','Relevant point','URL'],sources,font=7.5)

save_doc(d,'filing_ready_incorporation_and_board_execution_packet_v1_2.docx')

# ------------------ Doc 2: Director recruitment ------------------
d=base_doc('Founding Director Recruitment & Acceptance Kit','Recruit, screen, interview, appoint and onboard the additional people required for lawful and independent governance')
h1(d,'1. Immediate recruitment objective')
d.add_paragraph('The organization presently has one certain founder/director candidate. The legal minimum is three directors, but the governance target is five formation directors and later 7-11 directors. The immediate recruitment objective is therefore four additional formation directors, with at least two additional people required before lawful Board operation.')
callout(d,'Highest priority','Recruit an independent Chair candidate and a Secretary candidate first. The Secretary cannot be the same individual as the President. A finance-capable director and a safeguarding/integrity-capable director are the next priorities.')

table(d,['Seat','Preferred profile','Independence target','Priority'],[
['Founder / President','Founder; executive/religious/research leadership','Not independent','Existing'],
['Independent Chair','Governance judgment; willing to supervise and if necessary oppose founder/executive','Independent','1'],
['Secretary','Reliable records/minutes; governance discipline; separate from President','Independent preferred','1'],
['Finance / Treasurer director','Budget/accounting/controls; no dependence on founder','Independent','2'],
['Safeguarding / Integrity director','Child/vulnerable-person safety, investigations, ethics or compliance orientation','Independent','2'],
])

h1(d,'2. What a founding director is agreeing to govern')
d.add_paragraph('A director is not being asked to endorse every theological hypothesis or become a spiritual subordinate. A director is accepting fiduciary responsibility for a Tennessee nonprofit religious corporation and the duty to preserve lawful governance, mission, financial integrity, safeguarding, and institutional independence.')
bullets=[
'Protect the mission without protecting the founder from accountability.',
'Review finances, compensation, conflicts and major transactions rather than rubber-stamping them.',
'Ensure complaints against senior leaders can reach independent decision-makers.',
'Protect the constitutional distinction between religious community and corporate authority.',
'Ask difficult questions about private benefit, related-party transactions, housing/property, charitable assistance, children, fundraising and paid staff.',
'Be willing to vote against the founder or executive when fiduciary duty requires it.',
'Keep appropriate confidentiality while refusing secrecy that conceals abuse, illegality, fraud or safeguarding failures.',
]
for x in bullets: bullet(d,x)

h1(d,'3. Independence standard')
d.add_paragraph('A candidate may still be legally eligible while not being independent. For this project, independence is a governance category used to keep oversight credible.')
table(d,['Question','If yes, effect'],[
['Is the candidate the founder\'s spouse, parent, child, sibling, close relative, or intimate partner?','Normally not independent.'],
['Is the candidate employed by, paid by, housed by, or materially financially dependent on the founder or organization?','Normally not independent.'],
['Does the candidate own or control a vendor, landlord, lender, contractor or major counterparty?','Conflict; likely not independent for affected matters.'],
['Does the candidate have a major business partnership, debt, investment or shared property with the founder?','Potential loss of independence; disclose and assess.'],
['Does the candidate receive pastoral/counseling support from the founder in a way that may impair oversight?','Assess carefully; may impair independence.'],
['Would the candidate realistically investigate, suspend, or remove the founder if evidence required it?','A "no" is disqualifying for an independent oversight seat.'],
])

h1(d,'4. Candidate invitation - reusable draft')
d.add_paragraph('Suggested invitation language:')
callout(d,'Invitation','I am forming a Tennessee nonprofit religious organization around the Way, a developing religious community designed with unusually strong safeguards for conscience, safe exit, financial transparency, child protection, independent complaints, and limits on founder power. I am looking for founding directors who are willing to govern the institution rather than simply support me personally. The role includes fiduciary responsibility, review of finances and conflicts, oversight of executive leadership, and willingness to challenge or investigate the founder if necessary. You would receive the governing documents before deciding, and no theological loyalty to me is expected. Would you be willing to review the board information packet and have a candid conversation about whether the role fits you?')

h1(d,'5. Candidate interview guide')
questions=[
'What would make you vote against the founder or President even if you personally trusted him?',
'If a complaint alleged misconduct by the founder, what process would you expect before deciding anything?',
'How would you distinguish religious disagreement from misconduct?',
'What financial information would you expect to see regularly as a director?',
'How should the organization handle a transaction involving a founder-relative, director-owned company, or property controlled by an insider?',
'What would concern you about a religious organization providing employment, housing, education, pastoral care and spiritual authority to the same person?',
'If attendance or donations declined because leadership disclosed a serious problem honestly, how would you think about that tradeoff?',
'What should happen when a participant says a religious teaching harmed them but leadership believes the teaching is sincere?',
'Are there circumstances in which you would resign rather than support a Board decision? What are they?',
'Can you commit to reading materials, attending meetings, maintaining confidentiality appropriately, and completing annual conflict/independence disclosures?'
]
for i,q in enumerate(questions,1): num(d,f'{i}. {q}')

h1(d,'6. Candidate conflict and independence disclosure')
fields=['Full legal name','Home city/state','Current employer / business','Relationship to founder','Relationship to other proposed directors','Employment/consulting with organization','Vendor/landlord/lender relationship','Loans/debts/guarantees/shared ventures','Gifts/housing/financial support','Family/intimate relationships relevant to oversight','Pastoral/counseling dependency relevant to oversight','Litigation/material disputes','Other nonprofit/business offices','Other actual/apparent conflicts']
table(d,['Disclosure item','Candidate response'],[[x,'[[RESPONSE]]'] for x in fields],font=7.8)
d.add_paragraph('Candidate certification: I have disclosed relationships that could reasonably affect or appear to affect my judgment. I will update this disclosure when circumstances change and will follow recusal decisions. Signature: __________________ Date: __________')

h1(d,'7. Director acceptance and fiduciary acknowledgement')
d.add_paragraph('I, [[NAME]], accept appointment/election as a director of [[LEGAL_NAME]], Inc., effective [[DATE]], subject to successful legal formation and the governing documents. I understand that my duties run to the corporation and its exempt purposes, not to the founder personally. I agree to act in good faith, with appropriate care, in the organization\'s best interests, to disclose conflicts, to preserve safeguarding and non-retaliation systems, and to participate in independent review of insider matters.')
d.add_paragraph('Signature: ______________________________   Date: __________________')

h1(d,'8. Officer-specific role briefs')
roles=[
('Independent Chair','Organizes Board work, protects executive sessions, ensures founder/CEO evaluation, controls agenda access for governance matters, and supports independent investigations. Does not become the spiritual superior of the Assembly.'),
('Secretary','Maintains or oversees minutes and authenticated corporate records, tracks resolutions and terms, and ensures Board actions are actually documented. Must be a different person from the President.'),
('Treasurer / Finance lead','Oversees financial reporting, budget discipline, reconciliations, restricted funds, insurance/tax coordination and financial controls. Does not personally perform every bookkeeping function.'),
('Safeguarding / Integrity lead','Oversees safeguarding architecture, independent complaint routes, retaliation review, senior-leader conflict screening, external investigation triggers and integrity reporting.'),
]
for a,b in roles: h2(d,a); d.add_paragraph(b)

d.add_page_break()
h1(d,'9. Candidate decision matrix')
table(d,['Candidate','Independent?','Chair fit','Secretary fit','Finance fit','Safeguarding fit','Concerns','Next step'],[
['[[NAME 1]]','','','','','','',''],['[[NAME 2]]','','','','','','',''],['[[NAME 3]]','','','','','','',''],['[[NAME 4]]','','','','','','',''],['[[NAME 5]]','','','','','','',''],
],font=7.0)

h1(d,'10. First-30-day onboarding')
for i,x in enumerate([
'Read Charter draft, Bylaws, Constitution of the Way v1.0, Teaching & Practice Baseline, Companion Handbook and Integrity Framework.',
'Complete conflict/independence disclosure and identify recusals.',
'Review director duties, records, meeting cadence, insurance plan and document hierarchy.',
'Review founder compensation rule and related-party transaction process.',
'Review safeguarding/mandatory-reporting route and complaint/investigation protocol.',
'Complete tabletop exercises: founder complaint; child-safety disclosure; related-party transaction; financial-control failure.',
'Confirm officer/committee assignments and Board calendar.',
],1): num(d,f'{i}. {x}')

h1(d,'11. Recruitment boundaries')
for x in [
'Do not recruit directors merely because they are personally loyal, spiritually impressed, financially dependent, or unable to say no.',
'Do not promise future employment, housing, business, spiritual rank, or compensation in exchange for Board service.',
'Do not conceal the fact that directors may need to investigate or remove the founder/executive.',
'Do not describe a candidate as a director publicly until the lawful appointment/election is effective.',
'Do not allow urgency around incorporation to substitute for basic competence, independence, and character screening.'
]: bullet(d,x)

save_doc(d,'founding_director_recruitment_and_acceptance_kit_v1_2.docx')

# ------------------ Doc 3: Federal tax/banking ------------------
d=base_doc('EIN, Banking, Federal Exemption & Compliance Execution Packet','Post-formation tax identification, account opening, Form 1023 readiness, public-support planning and first-year compliance')
h1(d,'1. Order of operations')
for i,x in enumerate([
'File and obtain acceptance of the Tennessee nonprofit charter.',
'Complete incorporator/initial Board organization and officer appointments.',
'Apply for the EIN only after legal formation. The IRS expressly warns exempt organizations not to apply before legal formation.',
'Open bank/accounting infrastructure under Board authority and the EIN; do not use personal accounts for organizational funds.',
'Finalize three-year financial projections and Form 1023 activity narratives using actual planned operations.',
'Complete Board review of compensation, conflicts, related-party arrangements, intellectual property, charitable-assistance methods and public-charity classification request.',
'Electronically file the full Form 1023 through Pay.gov with the then-current user fee.',
'Operate under applicable filing, disclosure, charitable-solicitation, payroll and state tax rules while the determination is pending and afterward.'
],1): num(d,f'{i}. {x}')

h1(d,'2. EIN application worksheet')
d.add_paragraph('The IRS currently provides a free online EIN application and states that a legal entity should be formed with the state first. For tax-exempt organizations, the responsible party is generally the principal officer and must generally be a natural person with the required taxpayer identification number.')
table(d,['EIN field','Working entry'],[
['Legal name','[[LEGAL_NAME]], Inc. - exactly as filed with Tennessee'],
['Trade/DBA name','[[PUBLIC_NAME if different; otherwise blank]]'],
['Mailing address','[[MAILING_ADDRESS]]'],
['Physical address','[[PHYSICAL/PRINCIPAL_ADDRESS]]'],
['County/state','[[COUNTY]], Tennessee'],
['Responsible party','[[PRINCIPAL_OFFICER - likely President]]'],
['Entity type','Corporation / tax-exempt organization; select the IRS category that truthfully matches the legal entity and current operations.'],
['Reason applying','Started a new organization / banking / tax administration as applicable.'],
['Employees expected','[[GOOD-FAITH ESTIMATE]]'],
['Closing month','December'],
['Principal activity','Religious, charitable and educational activities.'],
])
callout(d,'Privacy','Do not place a Social Security number or other taxpayer identification number into shared governance drafts. Enter it only in the official IRS application when the responsible party submits the EIN request.')

h1(d,'3. Bank-account opening packet')
d.add_paragraph('Typical bank requirements vary, but the organizational file should be prepared with the following materials. Confirm the chosen institution\'s current requirements before the appointment.')
for x in ['Filed/accepted Tennessee charter','EIN confirmation','Bylaws','Board banking resolution','Officer/director identification as requested by bank','Registered/principal address information','Beneficial-owner/control-person certifications if the bank requests them under its customer-identification rules','Initial deposit source documentation','Signer and transaction-authority matrix']:
    bullet(d,x)

h2(d,'3.1 Model banking resolution')
d.add_paragraph('RESOLVED, that the Corporation is authorized to establish deposit, payment-processing, merchant and other financial accounts reasonably necessary for its exempt activities at institutions approved by the Board or its Finance Committee; provided that no account shall be titled in the name of an individual, organizational funds shall not be commingled with personal funds, and authority shall be limited by the approved signer/transaction matrix.')

table(d,['Transaction','Proposed control'],[
['Routine operating payments','Prepared/entered by authorized staff or officer; approval thresholds documented.'],
['Bank reconciliation','Performed/reviewed by a person who does not control the entire receipt-to-disbursement cycle where practicable.'],
['Large or unusual payments','Second approval above Board-set threshold.'],
['Executive/founder reimbursement','Independent approval and documented business purpose.'],
['Related-party/vendor payments','Conflict disclosure, recusal, alternatives/comparability and disinterested approval.'],
['Cash handling','Minimize cash; dual count/log where used; prompt deposit.'],
['Restricted gifts','Separate accounting by restriction; no diversion to unrestricted use without lawful basis.'],
])

h1(d,'4. Form 1023 execution strategy')
d.add_paragraph('The current strategy remains a full Form 1023 rather than Form 1023-EZ. Form 1023 must be filed electronically through Pay.gov. The IRS requires an EIN before Form 1023 can be filed. The organizing document must contain the required section 501(c)(3) limitations, and the application requires detailed activity information sufficient to evaluate public rather than private operation.')

table(d,['Form 1023 workstream','Required evidence before submission'],[
['Organizational documents','Filed charter; adopted bylaws; amendments if any.'],
['Officers/directors','Names, titles, addresses as required in the form; relationships and compensation information.'],
['Narrative activities','What; who conducts; when; where; how it furthers exempt purposes; approximate time allocation; how funded.'],
['Financial data','Good-faith current and projected revenue/expense information for the required periods.'],
['Compensation','Duties, total compensation, comparability data, independent approval process, conflicts/recusals.'],
['Insider transactions','Founder IP, leases, loans, reimbursements, vendors, property arrangements and any other private-party dealings.'],
['Fundraising','Methods, online giving, grants, events, paid fundraisers if any, donor restrictions and acknowledgments.'],
['Charitable assistance','Eligibility class, objective need/charitable criteria, records, independent approvals for insiders.'],
['Religious operations','Gatherings, formation, Houses of Prayer, ministry, publications, research, worship/prayer, pastoral care.'],
['Public-charity classification','Initial request presently anticipated under §509(a)(1)/§170(b)(1)(A)(vi), subject to actual financial support plan; church classification deferred until actual facts support it.'],
])

h2(d,'4.1 Church classification decision rule')
d.add_paragraph('Do not claim church classification merely because the organization is religious. The IRS currently uses a facts-and-circumstances approach and generally expects a congregation or other religious membership group. The instructions expressly state that an organization without an established congregation may be a religious organization that does not yet qualify as a church and may request church classification later.')

h1(d,'5. Public-support planning')
d.add_paragraph('The working initial public-charity route is §509(a)(1)/§170(b)(1)(A)(vi), which generally depends on broad public/government/public-charity support and the applicable public-support tests. The three-year projection workbook should be updated only with good-faith assumptions.')
for x in [
'Diversify contributions rather than relying on one founder or a few insiders where feasible.',
'Record grants and contribution sources in a way that supports later public-support calculations.',
'Avoid promising donor control over doctrine, employment, complaints, investigations or charitable beneficiaries.',
'Monitor large contributors because public-support calculations can treat portions of unusually large private contributions differently.',
'Review public-support status annually with tax counsel/accountant once actual revenue exists.'
]: bullet(d,x)

h1(d,'6. Contribution receipts and donor communications')
d.add_paragraph('Before public fundraising, establish written donor-receipt and restricted-gift procedures. The organization should never promise deductibility beyond what federal law supports, especially before an IRS determination letter is received if the organization is choosing to seek one.')
table(d,['Scenario','Communication/control'],[
['Ordinary donation','Provide accurate contemporaneous acknowledgment when required; state legal entity receiving gift.'],
['Restricted donation','Accept only restrictions the organization can lawfully honor; record restriction in writing.'],
['Goods/services provided','Use required quid-pro-quo disclosure rules where applicable.'],
['Donation before IRS determination','Use careful wording about pending application/anticipated exemption; obtain professional review for public fundraising copy.'],
['Gift to specific individual','Do not convert a donor-directed personal gift into a purported charitable contribution; charitable aid must remain under organizational discretion for a charitable class/purpose.'],
])

h1(d,'7. Tennessee charitable solicitation and annual corporate compliance')
d.add_paragraph('Tennessee currently identifies bona fide religious institutions as exempt from charitable-solicitation registration under its charitable solicitations framework, and it also identifies certain small organizations as exempt subject to annual exemption-request requirements. The organization should document the exemption actually relied upon and recheck requirements before public solicitation. Separately, Tennessee nonprofit corporations have annual-report obligations; the Secretary of State currently lists a $20 corporation annual-report fee and requires nonprofit annual reports to list a President and Secretary.')

table(d,['Compliance item','First-year action'],[
['Charitable solicitation','Determine and document whether bona fide religious-institution exemption applies; file any exemption request/documentation required by current Tennessee practice.'],
['Tennessee annual report','Calendar due date after formation; maintain President, Secretary, directors, registered agent/office and required report information.'],
['IRS annual return/notice','Determine Form 990/990-EZ/990-N obligation based on actual classification and receipts; church exceptions apply only if the organization actually qualifies.'],
['Payroll','Before paying employees, establish payroll withholding/reporting and workers\' compensation/employment compliance as applicable.'],
['Sales/use/property tax','Federal §501(c)(3) status does not automatically establish every Tennessee state/local tax exemption; apply separately where needed.'],
['Insurance','Evaluate general liability, directors/officers, abuse/molestation, property, workers\' compensation, cyber and other coverage appropriate to actual programs.'],
])

h1(d,'8. Form 1023 pre-submission Board certification')
checks=[
'The filed charter and adopted bylaws match the copies attached/submitted.',
'The activity narrative describes planned activities truthfully and specifically.',
'The three-year financials are good-faith projections and reconcile to planned activities.',
'All directors/officers and family/business relationships are disclosed as required.',
'Founder compensation is either zero or prospectively approved through an independent process supported by comparability data.',
'Founder intellectual property, leases, loans or other insider transactions are documented and reviewed.',
'Public-charity classification request matches the actual funding model.',
'No church classification is claimed unless current facts support it.',
'No representation in the application contradicts the Constitution of the Way, safeguarding policies or anti-coercion controls.',
'An authorized principal officer/director is prepared to digitally sign under penalties of perjury.'
]
for i,x in enumerate(checks,1): num(d,f'{i}. [ ] {x}')

h1(d,'9. Current federal/state source notes')
sources=[
('IRS - Get an EIN','Form state legal entity before applying; EIN is free; responsible party requirements.','https://www.irs.gov/businesses/small-businesses-self-employed/get-an-employer-identification-number'),
('IRS - EIN for exempt organization','Do not apply until legally formed.','https://www.irs.gov/charities-non-profits/obtaining-an-employer-identification-number-for-an-exempt-organization'),
('IRS - Form 1023','Electronic filing; EIN required; §501(c)(3) organizational test and public charity classification.','https://www.irs.gov/instructions/i1023'),
('IRS - Form 1023 FAQ','Activity narrative asks what/who/when/where/purpose/time/funding.','https://www.irs.gov/charities-non-profits/frequently-asked-questions-about-form-1023'),
('Tennessee Secretary of State','Current nonprofit forms/fees and annual-report guidance.','https://sos.tn.gov/businesses/forms-and-fees'),
('Tennessee charitable solicitation FAQ','Religious-institution and other exemptions.','https://sos.tn.gov/charitable/charitable-organizations'),
]
table(d,['Source','Use','URL'],sources,font=7.4)

save_doc(d,'ein_banking_federal_exemption_and_compliance_execution_packet_v1_2.docx')

# ------------------ Supporting source files ------------------
readiness=[
['Gate','Owner','Required before','Status','Evidence'],
['Legal name','Founder + counsel','State filing','OPEN','Name clearance record'],
['Minimum directors','Founder','State organization','OPEN','At least 3 names/acceptances'],
['Secretary','Board candidate','Organizational meeting','OPEN','Acceptance by person other than President'],
['Registered office','Founder','State filing','OPEN','Tennessee street address'],
['Principal office','Founder','State filing/EIN','OPEN','Physical/mailing address'],
['Counsel charter review','Founder/counsel','State filing','PENDING','Reviewed final charter'],
['Tennessee filing','Founder/incorporator','EIN','PENDING','Accepted charter'],
['Board organization','Initial board','Bank/Form 1023','PENDING','Minutes/consent'],
['EIN','Principal officer','Bank/Form 1023','PENDING','CP575/confirmation'],
['Bank/accounting','Treasurer/board','Public fundraising','PENDING','Account + controls'],
['Independent complaint route','Integrity lead','Pilot/public launch','OPEN','Published contact/process'],
['3-year projections','Board/treasurer','Form 1023','DRAFT','Completed workbook'],
['Form 1023 final review','Board','Submission','PENDING','Board certification'],
]
with open(SRC/'filing_readiness_register_v1_2.csv','w',newline='',encoding='utf-8') as f: csv.writer(f).writerows(readiness)

(Path(SRC/'director_candidate_tracker_v1_2.csv')).write_text('Candidate,Relationship to founder,Independent Y/N,Chair fit,Secretary fit,Finance fit,Safeguarding fit,Contacted,Interviewed,Disclosure complete,Decision,Notes\n',encoding='utf-8')

(Path(SRC/'name_clearance_record_v1_2.md')).write_text('''# Legal/Public Name Clearance Record v1.2\n\nUse for each serious candidate. No name is approved by this template itself.\n\n## Candidate\n- Legal corporate name: [[...]]\n- Public/community name: [[...]]\n- Domain/social variants: [[...]]\n\n## Checks\n- [ ] Tennessee business-name search / filing availability\n- [ ] Tennessee assumed-name conflicts if relevant\n- [ ] USPTO trademark search for confusingly similar religious/nonprofit/education marks\n- [ ] General web search for religious organizations, ministries, high-control/cult associations, scandals, extremist associations, and confusingly similar movements\n- [ ] Domain availability and obvious typos\n- [ ] Major social handles where relevant\n- [ ] Linguistic/cultural screening in expected languages\n- [ ] Counsel review if the name will carry significant brand value\n\n## Decision\nStatus: OPEN / HOLD / CLEAR FOR FILING SUBJECT TO COUNSEL\nRationale: [[...]]\nChecked by: [[...]]\nDate: [[...]]\n''',encoding='utf-8')

(Path(SRC/'execution_sequence_v1_2.md')).write_text('''# v1.2 Execution Sequence\n\n1. Select and clear the legal/public name.\n2. Recruit at least two additional directors; target four additional directors.\n3. Identify a Secretary different from the President.\n4. Complete conflict/independence disclosures.\n5. Insert registered/principal office information privately into final filing copy.\n6. Obtain final Tennessee nonprofit counsel review of Charter/Bylaws.\n7. File Tennessee charter and retain accepted/stamped copy.\n8. Complete incorporator action and initial Board organization.\n9. Apply for EIN only after legal formation.\n10. Open bank/accounting/payment infrastructure under Board resolution.\n11. Activate independent complaint/safeguarding routes before public launch.\n12. Finalize three-year financial assumptions and full Form 1023 application.\n13. Board reviews and authorizes Form 1023 filing.\n14. File electronically through Pay.gov.\n15. Begin/continue controlled pilot and first-year compliance calendar.\n''',encoding='utf-8')

meta={
 'version':'1.2','date':'2026-09-21','working_legal_name':'UNDECIDED Inc.','working_public_name':'UNDECIDED','corporate_model':'Tennessee public benefit religious nonprofit; no statutory members at formation (working decision)','fiscal_year_end':'December 31','founder_role':'Founder; intended President/Executive Director subject to Board appointment','current_directors':['Thomas Carter'],
 'open_gates':['legal name','2+ additional directors','Secretary other than President','registered office','principal office','final counsel review'],
 'documents':['filing_ready_incorporation_and_board_execution_packet_v1_2.docx','founding_director_recruitment_and_acceptance_kit_v1_2.docx','ein_banking_federal_exemption_and_compliance_execution_packet_v1_2.docx']
}
(SRC/'organizational_instantiation_v1_2.json').write_text(json.dumps(meta,indent=2),encoding='utf-8')

(Path(OUT/'README_RELEASE.md')).write_text('''# The Way v1.2 - Filing-Ready Incorporation & Founding Board Execution\n\nThis release converts the post-v1.0 religious system into a near-executable Tennessee corporate/federal tax formation package without fabricating unresolved legal name, people, addresses, or financial facts.\n\n## Publications\n- Filing-Ready Incorporation & Founding Board Execution Packet\n- Founding Director Recruitment & Acceptance Kit\n- EIN, Banking, Federal Exemption & Compliance Execution Packet\n\n## Controlled gates that still require real-world input\n1. Final clearable legal/public name.\n2. At least two additional natural-person directors; target four additional formation directors.\n3. Secretary different from President.\n4. Registered/principal office addresses.\n5. Final attorney review before Tennessee filing.\n6. Good-faith financial projections before Form 1023 filing.\n\nThe packet is structured so these values can be inserted without another governance-design pass.\n''',encoding='utf-8')

print('built')
