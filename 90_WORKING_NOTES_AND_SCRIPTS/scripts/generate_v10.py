from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import json, csv, textwrap

OUT = Path('/mnt/data/way_v10')
SRC = OUT/'source'

DARK = '1F2937'
MID = '475569'
LIGHT = 'E2E8F0'
PALE = 'F8FAFC'
ACCENT = '334155'
WHITE = 'FFFFFF'
RED = '991B1B'
GREEN = '166534'
AMBER = '92400E'


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_text(cell, text, bold=False, color=None, size=9):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP


def repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def set_repeat_table_headers(table):
    if table.rows:
        repeat_table_header(table.rows[0])


def set_keep_with_next(paragraph, value=True):
    pPr = paragraph._p.get_or_add_pPr()
    keepNext = pPr.find(qn('w:keepNext'))
    if value and keepNext is None:
        keepNext = OxmlElement('w:keepNext')
        pPr.append(keepNext)


def set_keep_lines(paragraph, value=True):
    pPr = paragraph._p.get_or_add_pPr()
    keepLines = pPr.find(qn('w:keepLines'))
    if value and keepLines is None:
        keepLines = OxmlElement('w:keepLines')
        pPr.append(keepLines)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = ' PAGE '
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)


def set_doc_defaults(doc, short_title):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal.font.size = Pt(10.2)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.06
    for sname, size, color in [('Title', 27, DARK), ('Heading 1', 17, DARK), ('Heading 2', 13, ACCENT), ('Heading 3', 11, MID)]:
        st = styles[sname]
        st.font.name = 'Aptos Display' if sname in ('Title','Heading 1') else 'Aptos'
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        if sname != 'Title':
            st.font.bold = True
            st.paragraph_format.space_before = Pt(9)
            st.paragraph_format.space_after = Pt(4)
    # header / footer
    hp = sec.header.paragraphs[0]
    hp.text = f'UNDECIDED | The Way | {short_title} v1.0'
    hp.style = styles['Normal']
    hp.runs[0].font.size = Pt(8)
    hp.runs[0].font.color.rgb = RGBColor.from_string(MID)
    fp = sec.footer.paragraphs[0]
    add_page_number(fp)
    for r in fp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor.from_string(MID)
    return doc


def title_page(doc, title, subtitle, status='Founding synthesis - working religious baseline'):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(90)
    r = p.add_run('THE WAY')
    r.bold = True; r.font.size = Pt(14); r.font.color.rgb = RGBColor.from_string(MID)
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(title)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(subtitle)
    r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string(MID)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run('Version 1.0 | September 20, 2026')
    r.bold = True; r.font.size = Pt(11)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(status)
    r.font.size = Pt(9.5); r.font.color.rgb = RGBColor.from_string(MID)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(55)
    r = p.add_run('Truth precedes the institution. Conscience is never surrendered. Authority is entrusted and limited.')
    r.italic = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor.from_string(ACCENT)
    doc.add_page_break()


def add_h1(doc, text):
    p = doc.add_paragraph(text, style='Heading 1'); set_keep_with_next(p); return p

def add_h2(doc, text):
    p = doc.add_paragraph(text, style='Heading 2'); set_keep_with_next(p); return p

def add_h3(doc, text):
    p = doc.add_paragraph(text, style='Heading 3'); set_keep_with_next(p); return p

def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); r.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    set_keep_lines(p)
    return p

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.add_run('• ').bold = True
        p.add_run(item)
        set_keep_lines(p)

def add_numbered(doc, items):
    for i, item in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.18)
        p.add_run(f'{i}. ').bold = True
        p.add_run(item)
        set_keep_lines(p)

def add_callout(doc, label, text, fill=PALE, color=DARK):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    repeat_table_header(t.rows[0])
    c = t.cell(0,0); shade_cell(c, fill)
    p = c.paragraphs[0]
    r = p.add_run(label + ': '); r.bold = True; r.font.color.rgb = RGBColor.from_string(color)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None, font_size=8.7):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    for j,h in enumerate(headers):
        c=table.rows[0].cells[j]; shade_cell(c, ACCENT); set_cell_text(c,h,True,WHITE,font_size)
    for row in rows:
        cells = table.add_row().cells
        for j,val in enumerate(row):
            set_cell_text(cells[j], val, False, None, font_size)
    set_repeat_table_headers(table)
    if widths:
        for row in table.rows:
            for j,w in enumerate(widths):
                row.cells[j].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table

# ---------------- Constitution ----------------
constitution_sections = []

doc = set_doc_defaults(Document(), 'Constitution of the Way')
title_page(doc, 'Constitution of the Way', 'Religious Constitution, Rights Charter, Community Order & Development Rule')

add_h1(doc, 'Preamble')
add_para(doc, 'We affirm that reality and truth precede every institution, teacher, text, council, generation, and interpretation. The Way exists to help human beings seek, discern, practice, and transmit what is true and good without surrendering conscience to human authority.')
add_para(doc, 'We therefore constitute the Way as a living religious tradition walked by Companions, gathered as the Assembly, and expressed locally through Houses of Prayer. We receive the religious inheritance of humanity as a Library of Witnesses to be studied with reverence, historical honesty, reason, prayer, experience, and disciplined discernment. We do not assume that all witnesses say the same thing, that every inherited claim is true, or that uncertainty is disloyalty.')
add_para(doc, 'This Constitution is a religious constitution. It does not replace the charter, bylaws, fiduciary duties, employment rules, safeguarding obligations, or other civil-law instruments of any nonprofit corporation serving the Way. Where civil law assigns authority to a legal body, spiritual language does not silently transfer that authority elsewhere.')
add_callout(doc, 'Founding maxim', 'The institution exists to serve the Way; the Way does not exist to preserve the institution.')

add_h1(doc, 'Article I - Nature, Purpose, and Scope')
add_h2(doc, '1.1 The Way')
add_para(doc, 'The Way is a living religious path concerned with humanity\'s relationship to ultimate reality, truth, and the manner in which human life ought to be lived in response to that reality. It is practiced rather than merely subscribed to.')
add_h2(doc, '1.2 Purpose')
add_bullets(doc, [
    'To seek truth concerning ultimate reality, religious encounter, human flourishing, suffering, moral responsibility, and the ends of human life.',
    'To cultivate prayer, attention, study, discernment, ethical practice, service, companionship, hospitality, repair, and reverence.',
    'To preserve and examine historical and living religious testimony without flattening differences among traditions.',
    'To form communities in which questioning, disagreement, outside learning, criticism, and departure can occur without retaliation.',
    'To transmit the Way across generations without making any founder, leader, office, or institution indispensable.'
])
add_h2(doc, '1.3 What the Way is not')
add_bullets(doc, [
    'The Way is not identical with a corporation, board, building, founder, minister, canon, creed, or generation.',
    'The Way is not a claim that all religions are secretly identical.',
    'The Way is not a license for leaders to convert private spiritual certainty into coercive authority.',
    'The Way is not complete merely because this Constitution reaches version 1.0.'
])

add_h1(doc, 'Article II - Foundational Constitutional Principles')
principles = [
    ('Reality precedes institution', 'No body can make a metaphysical claim true by vote, repetition, office, wealth, or age.'),
    ('Conscience is inalienable', 'No Companion surrenders moral or spiritual conscience to founder, Minister, Council, House, Assembly, or corporation.'),
    ('Participation is voluntary', 'Belonging, prayer, offerings, ministry, residence, employment, rites, study, and departure must remain free from coercion.'),
    ('Authority is functional and limited', 'Every office has a defined purpose, scope, review process, and limits.'),
    ('Evidence and interpretation are distinct', 'Encounter, Testimony, Interpretation, Teaching, and Doctrine must not be collapsed into one category.'),
    ('Correction is fidelity', 'New evidence, better argument, recognized harm, or deeper understanding may require revision of inherited teaching.'),
    ('Human beings are fallible', 'Governance must assume bias, self-deception, conflicts of interest, and misuse of power are possible in every person.'),
    ('The vulnerable are protected before reputation', 'Safeguarding, lawful reporting, immediate safety, and independent review take priority over institutional image or leader status.'),
    ('Exit is protected', 'A person can leave without shunning, retaliation, loss of ordinary relationships, or misuse of confidential information.'),
    ('No founder supremacy', 'Founder status is historical only and confers no doctrinal infallibility, veto, succession right, financial entitlement, or immunity.')
]
add_table(doc, ['Principle','Constitutional meaning'], principles, widths=[2.1,4.9])

add_h1(doc, 'Article III - Native Identity and Language')
add_table(doc, ['Layer','Native term','Meaning'], [
    ('Path / tradition','the Way','The religious way of life and inquiry.'),
    ('Person','Companion','One who freely walks the Way with others.'),
    ('People','the Companions','Companions considered collectively; not a spiritual caste.'),
    ('Religious community','the Assembly','The whole gathered religious community; distinct from a legal board.'),
    ('Local community','House of Prayer','A stable local expression of the Assembly; community before building.'),
    ('Ceremonial expression','House of Prayer for All Peoples','A fuller expression of hospitality and universality.'),
    ('Meeting','Gathering','A particular meeting for prayer, study, service, fellowship, or discernment.'),
], widths=[1.6,2.1,3.3])
add_para(doc, 'A Companion is not automatically a statutory corporate member. The Assembly is not automatically the nonprofit corporation. A House of Prayer is not automatically a separate legal entity. Religious vocabulary may not be used to obscure civil authority, ownership, employment, or financial responsibility.')

add_h1(doc, 'Article IV - Rights of Persons and Companions')
rights = [
    'Freedom to believe, doubt, question, disagree, reinterpret, or decline a religious claim.',
    'Freedom to read critics, consult other traditions, seek independent legal/medical/financial/mental-health advice, and maintain outside relationships.',
    'Freedom to attend public gatherings without becoming a Companion.',
    'Freedom to decline prayer forms, rites, fasting, counseling, donation requests, volunteer service, housing arrangements, or ministry roles.',
    'Freedom to leave immediately upon clear communication, without permission, spiritual release, confession, payment, nondisparagement agreement, or doctrinal recantation.',
    'Protection from shunning directives, public humiliation, doxxing, retaliatory disclosure, threats of supernatural catastrophe, or adverse action motivated by criticism or protected reporting.',
    'Access to safeguarding, complaint, independent investigation, appeal, and non-retaliation processes.',
    'Protection of children and vulnerable persons under applicable law and organizational safeguarding standards.',
    'Equal human dignity regardless of membership, office, wealth, sex, family status, certainty, disability, ethnicity, nationality, prior religion, criticism, or departure.'
]
add_bullets(doc, rights)
add_callout(doc, 'Entrenched right', 'No later teaching, prophecy, rite, council decision, House practice, or policy may silently cancel these rights.')

add_h1(doc, 'Article V - Companion Relationship and Commitment')
add_h2(doc, '5.1 Becoming a Companion')
add_para(doc, 'A person becomes a Companion only through a voluntary, informed decision after sufficient opportunity to understand the Way\'s practices, rights, current teachings, open questions, governance safeguards, and exit protections. No donation, property transfer, severing of outside relationships, employment arrangement, confession, or loyalty oath may be required.')
add_h2(doc, '5.2 Companion Commitment')
commitment = [
    'I choose freely to walk the Way in companionship with others.',
    'I will seek truth with humility and practice prayer, study, ethical action, service, and discernment according to conscience and capacity.',
    'I will distinguish testimony from interpretation, authority from truth, and conviction from coercive command.',
    'I will not surrender my conscience to a leader, office, council, or institution, nor demand that another person surrender conscience to me.',
    'I will respect the freedom of others to question, disagree, seek outside counsel, participate, or depart.',
    'I will support safeguarding, financial integrity, honest inquiry, accountability, and repair when harm occurs.',
    'I understand that I may cease identifying as a Companion at any time without permission, penalty, shunning, or loss of ordinary human relationships.'
]
add_bullets(doc, commitment)
add_h2(doc, '5.3 Belonging is not rank')
add_para(doc, 'No Companion is required to become a Minister, Teacher, Steward, Elder, donor, employee, resident, or officeholder. A person may remain an ordinary Companion for life without lower spiritual standing.')

add_h1(doc, 'Article VI - The Assembly')
add_para(doc, 'The Assembly is the religious community formed when Companions and welcome participants gather around the Way. Its functions include prayer, study, inquiry, service, fellowship, mutual care, teaching, formation, discernment, rites, remembrance, and transmission.')
add_bullets(doc, [
    'The Assembly may deliberate religious questions and express communal discernment.',
    'The Assembly may create councils with defined religious responsibilities.',
    'The Assembly may not use acclamation, prophecy, majority sentiment, or founder preference to suspend participant rights, safeguarding, financial controls, due process, or lawful fiduciary duties.',
    'Religious consensus does not automatically bind a civil corporation unless adopted through the corporation\'s lawful processes.'
])

add_h1(doc, 'Article VII - Houses of Prayer')
add_para(doc, 'A House of Prayer is a stable local expression of the Assembly. It is a community before it is a building and may gather in a home, rented room, public facility, dedicated property, or an appropriate online/hybrid form.')
add_table(doc, ['Stage','Meaning'], [
    ('Gathering Circle','Recurring group exploring the Way; no claim to official House status.'),
    ('House-in-Formation','Meets minimum practice, safeguarding, and governance requirements.'),
    ('Recognized House of Prayer','Recognized by written instrument defining scope, responsibilities, finances, complaint routes, and review.'),
    ('Mature House','Demonstrates stable practice, leadership rotation, transparent finances, service, formation, and healthy independence from one charismatic person.'),
    ('Dormant / Closed House','Recognition ends without shame when the community no longer functions; people remain free and property/records follow legal arrangements.')
], widths=[2.0,5.0])
add_h2(doc, '7.1 Minimum functions')
add_bullets(doc, ['Regular gathering','Prayer or worship','Study and inquiry','Service or mutual care','Safeguarding contact and independent escalation path','Local stewardship accountable to written authority'])
add_h2(doc, '7.2 House limitations')
add_para(doc, 'A House may not independently incur debt, employ staff, open accounts, sign property agreements, issue credentials, or claim separate legal personality without written civil authority. No House leader is a local sovereign.')

add_h1(doc, 'Article VIII - Ministry, Service, and Office')
add_table(doc, ['Role','Purpose','Constitutional limit'], [
    ('Teacher','Teach within demonstrated competence.','Subject to sources, review, correction, and scope.'),
    ('Minister','Provide defined prayer, pastoral, rite, teaching, or service functions.','Commissioned function, not sacred caste or immunity.'),
    ('Steward','Administer resources, operations, property, records, or local coordination.','Custody on behalf of community, not ownership.'),
    ('Elder / Mentor (provisional)','Recognize maturity, wisdom, service, and sound judgment.','No automatic executive, financial, disciplinary, or revelatory power.'),
    ('Founder','Historical originator of the present institutional project.','No inherent veto, infallibility, succession right, compensation right, or immunity.')
], widths=[1.5,2.8,2.7])
add_h2(doc, '8.1 Commissioning')
add_para(doc, 'Commissioning is the present authorization process for ministry. Every commission must state scope, supervision, safeguarding requirements, financial authority if any, term/review date where applicable, complaint route, and explicit limits. Commissioning is reviewable and revocable through fair process.')
add_h2(doc, '8.2 Ordination')
add_para(doc, 'Ordination remains reserved. It shall not be adopted merely to imitate another tradition, manufacture tax status, or create a permanent clerical class. If later adopted, its theological claims and institutional consequences require a separate dossier and constitutional review.')

add_h1(doc, 'Article IX - Epistemic Integrity and Teaching Authority')
add_h2(doc, '9.1 Required distinctions')
add_para(doc, 'The Way distinguishes Encounter -> Testimony -> Interpretation -> Teaching -> Doctrine. A powerful experience, sincere witness, revered text, founder statement, prediction, dream, vision, altered state, or claimed revelation does not bypass the interpretive steps between report and institutional teaching.')
add_h2(doc, '9.2 Confidence grammar')
status_rows = [
    ('Established','Strong evidence and broad responsible agreement for the defined claim; still revisable.'),
    ('Well-Supported','Substantial evidence favors the claim; meaningful uncertainty remains.'),
    ('Provisional','Plausible working conclusion useful for practice or inquiry and explicitly subject to revision.'),
    ('Permitted Interpretation','Coherent position a Companion may hold without Assembly endorsement as the best-established view.'),
    ('Open Question','Evidence does not justify an Assembly conclusion; no conformity required.'),
    ('Rejected as Harmful or Unsupported','Either evidentially untenable for the defined claim or impermissible as a governing use because it violates rights/safeguards; reasons must be stated.')
]
add_table(doc, ['Status','Meaning'], status_rows, widths=[1.8,5.2])
add_h2(doc, '9.3 Doctrine')
add_para(doc, 'Doctrine is a formally adopted teaching with defined status and implications. Doctrine must be used sparingly. No opinion becomes doctrine by repetition, popularity, founder usage, liturgical habit, donor pressure, or age.')
add_h2(doc, '9.4 What would change our mind')
add_para(doc, 'Every material theological proposal should identify the evidence, argument, recognized harm, or counterexample that would justify revision. A proposition defined so that no conceivable evidence can count against it may be held as faith or hope, but it must not be misrepresented as historical or empirical knowledge.')

add_h1(doc, 'Article X - Library of Witnesses')
add_para(doc, 'The Assembly maintains a Library of Witnesses rather than a closed canon at this founding stage. The Library may include scriptures, oral traditions, historical records, philosophical works, material artifacts, liturgies, testimonies, scholarship, practices, and contemporary reports relevant to the Way\'s inquiry.')
add_bullets(doc, [
    'Inclusion signifies relevance for study, not inspiration, infallibility, equal authority, or total endorsement.',
    'Source traditions should be represented through primary sources, practitioners, and responsible scholarship whenever possible.',
    'Historical provenance, transmission, genre, translation, competing interpretations, and known limitations should be recorded.',
    'No founder-authored text receives automatic canonical or superior status.',
    'A future canon, if any, requires a distinct theological process and may not be created by administrative convenience.'
])

add_h1(doc, 'Article XI - Practice of the Way')
add_h2(doc, '11.1 Sevenfold practice')
add_table(doc, ['Practice','Purpose'], [
    ('Attention','Cultivate presence, honesty, awareness, and freedom from automatic reaction.'),
    ('Prayer','Address, listen, praise, lament, petition, give thanks, remember, and dedicate oneself toward God/the Divine as honestly understood.'),
    ('Study','Learn from sacred sources, history, languages, scholarship, science, philosophy, and lived traditions.'),
    ('Discernment','Test claims, motives, choices, and interpretations individually and communally.'),
    ('Ethical Practice','Train honesty, restraint, justice, compassion, responsibility, fidelity to consent, and repair.'),
    ('Service','Act for the good of people, communities, creatures, and the world without conditioning aid on conversion or loyalty.'),
    ('Companionship','Walk with others through fellowship, mutual care, accountability, hospitality, and shared life.')
], widths=[1.5,5.5])
add_h2(doc, '11.2 Gatherings')
add_bullets(doc, ['Gathering for Prayer','Gathering for Study','Gathering for Service','Gathering for Fellowship / Common Table','Gathering for Discernment'])
add_h2(doc, '11.3 Rule of Life')
add_para(doc, 'A Rule of Life may guide regular prayer, study, ethical practice, service, rest, community participation, and self-examination. It must remain adaptable to health, disability, family responsibilities, work, culture, and conscience and must never become a surveillance or loyalty system.')

add_h1(doc, 'Article XII - Ethical and Pastoral Baseline')
ethics = [
    'Do not presume that illness, abuse, disability, bereavement, poverty, disaster, infertility, mental distress, or other suffering proves divine punishment, karmic desert, weak faith, or spiritual inferiority.',
    'Do not threaten damnation, rebirth into suffering, divine abandonment, loss of salvation, or spiritual catastrophe to obtain money, sex, labor, secrecy, obedience, housing compliance, votes, or continued participation.',
    'Capacity, meaningful consent, non-exploitation, honesty, safeguarding, and responsibility to freely undertaken commitments and dependents are minimum standards for intimate relationships.',
    'Violence, stalking, sexual assault, child abuse, grooming, and coercive control are safety/legal matters before pastoral reconciliation matters.',
    'The Way maintains a strong presumption toward nonviolence, peacemaking, de-escalation, and protection of vulnerable persons; absolute pacifism remains an open question.',
    'No religious actor may authorize vigilantism, holy war, coercive enforcement, or revelatory exemption from law.',
    'Pastoral care may accompany decisions but may not dictate another adult\'s healthcare, legal, financial, marital, reproductive, residential, educational, or employment choices.',
    'Forgiveness, reconciliation, restoration of relationship, and restoration to office are distinct. No one must resume contact or return a person to authority as proof of forgiveness.'
]
add_bullets(doc, ethics)

add_h1(doc, 'Article XIII - Sacred Time, Rites, and Hope')
add_h2(doc, '13.1 Rhythm before calendar')
add_bullets(doc, [
    'Daily: voluntary attention/prayer and honest self-examination.',
    'Weekly: recurring communal gathering for prayer, study, fellowship, service, or rest; no universal required weekday is yet constitutionalized.',
    'Periodic: Common Table, service, extended study, discernment, retreat, generosity, remembrance, or optional disciplines.',
    'Seasonal/annual: observances may be piloted and reviewed before becoming permanent sacred calendar obligations.'
])
add_h2(doc, '13.2 Rites')
add_para(doc, 'Rites may express commitment, transition, remembrance, healing, mourning, commissioning, reconciliation, and communal identity. Their language must distinguish symbolic action, inherited tradition, hope, provisional belief, and established claim. No rite may create an irreversible legal, sexual, financial, or spiritual transfer of control by ambiguity.')
add_h2(doc, '13.3 Hope without manufactured certainty')
add_para(doc, 'The Way may pray for healing, reconciliation, mercy, peace, liberation, awakening, and ultimate good, and may study resurrection, rebirth, paradise, judgment, nirvana, universal reconciliation, and other hopes. Pastoral need does not authorize factual certainty the evidence does not support.')

add_h1(doc, 'Article XIV - Interreligious Life')
add_bullets(doc, [
    'Learn a tradition first through its primary sources, practitioners, history, and responsible scholarship before criticizing it.',
    'Distinguish "this tradition teaches X" from "all adherents believe X" and from "the Way judges X true."',
    'Cooperate with other religious and secular communities in charity, peacebuilding, education, disaster relief, and common goods without pretending doctrinal agreement.',
    'Do not require Companions to sever family or friendships because of religious difference.',
    'Dual religious belonging remains an open theological question to be handled with case-specific honesty rather than blanket prohibition or blanket endorsement.',
    'No compelled conversion, deceptive recruitment, or appropriation of another tradition\'s holy day, title, rite, or vocabulary as though the Way originated it.'
])

add_h1(doc, 'Article XV - Religious Community and Civil Institution')
add_h2(doc, '15.1 Distinct bodies')
add_para(doc, 'The Assembly is the religious community. The nonprofit corporation is a civil legal instrument for property, contracts, employment, finance, risk management, and continuity. The Board of Directors governs the corporation under law and governing documents. Religious councils govern only the religious matters expressly entrusted to them.')
add_h2(doc, '15.2 Financial and institutional integrity')
add_bullets(doc, [
    'Giving is voluntary; money never purchases spiritual status, salvation, prophecy, access, office, healing, or preferential treatment.',
    'Founder and insider compensation requires independent, disinterested approval and appropriate comparability support.',
    'Related-party transactions require disclosure, recusal, fairness review, and documentation.',
    'Complaint, safeguarding, investigation, and financial controls remain operative even when allegations involve the founder or senior spiritual leadership.'
])

add_h1(doc, 'Article XVI - Founder, Succession, and Transmission')
add_para(doc, 'The founder is the historical originator of the present institutional project, not the permanent owner of the Way. The Way must be designed to survive the founder\'s death, incapacity, resignation, retirement, disagreement, or prolonged absence.')
add_bullets(doc, [
    'No successor inherits founder status.',
    'No founder statement is self-authenticating revelation or doctrine.',
    'No founder family relationship creates hereditary office.',
    'Archives should preserve founder reasoning, evidence, dissent, revisions, and mistakes rather than converting memory into hagiography.',
    'The founder-exit test is passed when strangers can discover the Way, become Companions, form healthy Houses, train accountable servants, and transmit the tradition without direct access to the founder.'
])

add_h1(doc, 'Article XVII - Development, Review, and Amendment')
add_h2(doc, '17.1 Protected core')
add_para(doc, 'The following principles are constitutionally entrenched and may not be reduced by ordinary teaching or policy: freedom of conscience; voluntary participation and exit; non-retaliation; safeguarding and lawful reporting; no infallible living authority; no founder supremacy; separation of revelation claim from coercive command; financial integrity; due process; and the distinction between religious community and civil fiduciary authority.')
add_h2(doc, '17.2 Ordinary development')
add_para(doc, 'Teaching, liturgy, formation, calendars, ministry forms, and local practices may develop through evidence, lived experience, discernment, pilot practice, rights-impact review, and recorded decision. Change should be easier where claims are provisional and harder where rights or constitutional safeguards are affected.')
add_h2(doc, '17.3 Research rule')
add_para(doc, 'Material doctrinal change requires a written record of the proposition, sources, arguments, competing interpretations, confidence level, practical implications, rights impact, dissent, decision process, effective date, and review date.')

add_h1(doc, 'Article XVIII - Continuity and the End of Institutions')
add_para(doc, 'A House of Prayer may close. A council may dissolve. A corporation may reorganize or cease. Buildings may be sold. Names and languages may change. None of these events is identical with the end of the Way.')
add_para(doc, 'The religious project reaches institutional maturity when it can faithfully reproduce its practices, inquiry, rights culture, and accountability across generations without dependence on any founder or single organization. Its ultimate spiritual end remains open to theological discovery rather than institutional invention.')
add_callout(doc, 'Completion test', 'A person can discover the Way, participate safely, become a Companion freely, mature, serve, question leaders, report misconduct, help form another House, transmit the tradition - or leave and criticize it - without requiring the founder and without suffering institutional retaliation.', fill='E8F5E9', color=GREEN)

add_h1(doc, 'Appendix A - Governing Hierarchy and Source Lineage')
add_table(doc, ['Layer','Function'], [
    ('Applicable civil law','Controls civil legal duties.'),
    ('Corporate charter and bylaws','Govern the nonprofit corporation and fiduciary authority.'),
    ('Institutional Integrity / Anti-Coercion protections','Operational safeguards for rights, safeguarding, complaints, finance, investigations, and founder limits.'),
    ('Constitution of the Way','Religious constitutional identity, rights, community order, epistemic and development principles.'),
    ('Teaching & Practice Baseline','Current teachings/practices with confidence status.'),
    ('House / ministry policies and liturgies','Local and functional implementation subject to higher protections.')
], widths=[2.4,4.6])
add_para(doc, 'Primary synthesis lineage: Native Religious Lexicon v0.5; Community Architecture v0.6; Theology, Practice & Formation Architecture v0.7; Provisional Theological Baseline v0.8; Provisional Ethical & Eschatological Baseline and operational framework v0.9; Institutional Integrity & Anti-Coercion governance system v0.2-v0.4.')

constitution_path = OUT/'constitution_of_the_way_v1_0.docx'
doc.save(constitution_path)

# ---------------- Teaching & Practice Baseline ----------------
doc = set_doc_defaults(Document(), 'Teaching & Practice Baseline')
title_page(doc, 'Teaching & Practice Baseline', 'What the Way presently teaches, practices, permits, and leaves open', 'Founding synthesis - not a creed of compulsory assent')

add_h1(doc, '1. Purpose and use')
add_para(doc, 'This baseline is the public operational teaching standard for the founding period. It does not demand assent to every proposition as a condition of human dignity, attendance, friendship, or charitable assistance. It tells Teachers, Ministers, Houses of Prayer, and Companions which claims may be taught directly, which require qualification, and which must remain open.')
add_callout(doc, 'Rule', 'Confidence may never outrun evidence, and practice may never outrun rights.')

add_h1(doc, '2. Confidence grammar')
add_table(doc, ['Status','Public teaching rule'], [
    ('Established','May be taught directly while grounds and revisability remain visible.'),
    ('Well-Supported','Teach as the best current conclusion with significant limitations acknowledged.'),
    ('Provisional','Teach as a working conclusion or experimental practice under review.'),
    ('Permitted Interpretation','May be held, explored, or used devotionally; no conformity required.'),
    ('Open Question','Present competing views fairly; the Assembly has no settled conclusion.'),
    ('Rejected as Harmful or Unsupported','May be studied historically, but may not be used to govern, coerce, or harm participants; reasons must be recorded.')
], widths=[1.8,5.2])

add_h1(doc, '3. Foundational teaching baseline')
core_teachings = [
    ('TB-001','Ultimate reality','Established','Religions make materially different claims about God/ultimate reality; the Way does not erase those differences.'),
    ('TB-002','God-language','Provisional','Companions and Houses may pray using ordinary God-language and other honest devotional language; no comprehensive metaphysics of God is yet fixed.'),
    ('TB-003','Revelation method','Established','A revelation claim is testimony requiring evaluation; intensity, sincerity, office, charisma, prediction, or altered state does not itself establish divine origin.'),
    ('TB-004','Revelation authority','Established','No claimed revelation overrides consent, safeguarding, law, financial controls, governing documents, or another adult\'s conscience.'),
    ('TB-005','Jesus - history','Established','Jesus was a first-century Jewish figure whose movement continued after his crucifixion under Pontius Pilate.'),
    ('TB-006','Jesus - metaphysics','Open Question','Incarnation, ontological divine sonship, Trinity, atonement mechanisms, miracle causation, and resurrection metaphysics are not presently Way doctrine.'),
    ('TB-007','Muhammad - history','Established','Muhammad was a seventh-century Arabian religious and political leader; the Qur\'an is the earliest principal textual witness to his proclamation.'),
    ('TB-008','Muhammad - revelation/finality','Open Question','Divine authorship of the Qur\'an and finality of prophethood are studied seriously but are not present Assembly doctrine.'),
    ('TB-009','Buddha - history/teaching','Established','A historical Gautama stands behind Buddhism; early teaching strongly centers suffering, causation, cessation, disciplined practice, impermanence, and non-self.'),
    ('TB-010','Buddha - metaphysics','Open Question','Rebirth, karma across lives, nirvana\'s metaphysical nature, and final implications of non-self remain open.'),
    ('TB-011','Scripture/canon','Established','The Assembly maintains a Library of Witnesses rather than a closed canon.'),
    ('TB-012','Human dignity','Established','Human dignity and conscience do not depend on membership, office, purity status, wealth, sex, belief certainty, or loyalty to the institution.'),
    ('TB-013','Human fallibility','Established','Founders, ministers, councils, scholars, and Companions remain vulnerable to bias, self-deception, harmful desire, social influence, and error.'),
    ('TB-014','Human transformation','Well-Supported','Attention, honesty, compassion, restraint, study, service, relationship, accountability, and repair are defensible formation goals across multiple anthropological models.'),
    ('TB-015','Soul / afterlife','Open Question','Immortal soul, bodily resurrection, rebirth, postmortem consciousness, annihilation, eternal separation, and universal reconciliation remain unresolved.'),
]
add_table(doc, ['ID','Domain','Status','Current baseline'], core_teachings, widths=[0.7,1.5,1.25,3.55], font_size=8.2)

add_h1(doc, '4. Ethical and pastoral baseline')
ethical_rows = [
    ('Suffering','Established','Do not presume suffering is deserved, divinely inflicted, karmically deserved, or evidence of weak faith; prioritize care, lament, treatment, justice, and safety.'),
    ('Theodicy','Open Question','No single explanation of suffering is adopted; proposed theodicies may not silence victims or excuse preventable harm.'),
    ('Spiritual end','Provisional','Healing, liberation, reconciliation, awakening, wholeness, and right relationship may be used as comparative/working terms without declaring them metaphysically identical.'),
    ('Exclusive salvation','Open Question','No current teaching that only Companions can be saved/liberated; no leader may claim knowledge of another person\'s ultimate destiny.'),
    ('Intimate ethics','Established','Capacity, meaningful consent, non-exploitation, honesty, safeguarding, and responsibility to freely undertaken commitments/dependents are minimum standards.'),
    ('Marriage theology','Open Question','Marriage sacramentality, same-sex blessing, divorce/remarriage, contraception, gender/sex metaphysics, plural marriage, celibacy, and related questions remain open.'),
    ('Violence','Well-Supported','Strong presumption toward nonviolence, peacemaking, de-escalation, and protection of vulnerable persons; absolute pacifism remains open.'),
    ('Religious violence','Established','No House, Minister, founder, Elder, or council may authorize vigilantism, holy war, coercive enforcement, or revelation-based exemption from law.'),
]
add_table(doc, ['Domain','Status','Operational baseline'], ethical_rows, widths=[1.6,1.4,4.0])

add_h1(doc, '5. Practice baseline')
add_h2(doc, '5.1 Daily and personal')
add_bullets(doc, [
    'Regular attention or stillness sufficient to cultivate honest awareness rather than dissociation or forced altered states.',
    'Prayer or contemplative practice according to conscience, using truthful language rather than compelled formulas.',
    'Study or reflection appropriate to season and capacity.',
    'Ethical self-examination focused on responsibility and repair rather than humiliation.',
    'Acts of care, generosity, service, or justice according to capacity.'
])
add_h2(doc, '5.2 Weekly and communal')
add_bullets(doc, [
    'At least one recurring House or Assembly gathering rhythm for prayer, study, fellowship, service, or rest.',
    'Regular exposure to primary sources and competing interpretations rather than only leader summaries.',
    'Opportunities for questions, dissent, testimony, and silence without pressure to perform spirituality.',
    'A safeguarding and complaint route that does not depend on the local spiritual leader.'
])
add_h2(doc, '5.3 Prayer grammar')
add_numbered(doc, ['Stillness / attention','Remembrance','Praise and gratitude','Truth and examination','Petition and intercession','Listening and discernment','Dedication and sending'])
add_para(doc, 'These movements are a grammar, not a mandatory script. Houses may develop language and music locally while preserving consent, theological honesty, accessibility, and the prohibition on binding personal revelations.')

add_h1(doc, '6. Gathering baseline')
add_table(doc, ['Gathering','Minimum purpose','Minimum integrity requirement'], [
    ('Prayer','Devotion, silence, gratitude, lament, intercession, remembrance.','No emotional manipulation, prophecy-as-command, forced disclosure, exhaustion, or restricted exit.'),
    ('Study','Disciplined learning and inquiry.','Use primary sources/context, competing interpretations, and explicit uncertainty.'),
    ('Service','Practical care.','No conversion, donation, loyalty, or silence requirement attached to aid.'),
    ('Fellowship / Common Table','Hospitality, equality, relationship.','No status seating or exclusion from ordinary food because of belief, wealth, criticism, or departure.'),
    ('Discernment','Test a question or proposed action.','State question, disclose conflicts, hear evidence/arguments, allow silence, record outcome and uncertainty.')
], widths=[1.6,2.6,2.8])

add_h1(doc, '7. Rites, sacred time, and fasting')
add_bullets(doc, [
    'Rites may express commitment, memory, transition, healing, mourning, commissioning, reconciliation, and community; metaphysical efficacy must be labeled at the confidence actually warranted.',
    'Companion Commitment and Commissioning are presently usable forms. Child welcome, healing prayer, mourning/memorial, and reconciliation may be used provisionally with safeguards. Marriage/covenant blessing and ordination remain reserved.',
    'Develop rhythm before comprehensive calendar: daily practice, weekly gathering, periodic service/Common Table/study, and seasonal/annual remembrance may be piloted and reviewed.',
    'Inherited holy days may be observed or studied according to conscience but must remain identified with their source traditions unless separately adopted with clear theological reason.',
    'Fasting is optional/experimental until a fuller theology and safety standard exists. Health, disability, pregnancy, eating disorders, medication, strenuous labor, age, and other risks require full accommodation without spiritual penalty.'
])

add_h1(doc, '8. Ministry baseline')
add_bullets(doc, [
    'Commissioning, not ordination, is the current authorization model.',
    'Every commission states scope, term/review, supervision, safeguarding requirements, complaint route, and financial authority if any.',
    'No commission creates ontological superiority, infallibility, immunity, sexual entitlement, automatic confidentiality privilege, or permanent office.',
    'Minister formation must include theology/religious literacy, source method, prayer leadership, pastoral boundaries, safeguarding, ethics/power, governance, conflict, referrals, practicum, and continuing review.',
    'A Minister may and should say "I do not know" where the baseline marks an Open Question.'
])

add_h1(doc, '9. Interreligious baseline')
add_bullets(doc, [
    'Represent traditions accurately and preserve real disagreement.',
    'Dialogue does not require relativism; conviction does not justify coercion.',
    'Cooperate in common good without pretending doctrinal agreement.',
    'Becoming a Companion does not require denouncing ancestry, relatives, former communities, or every prior belief.',
    'Dual religious belonging remains open and must be handled honestly where another tradition treats initiation, covenant, creed, or law as exclusive.'
])

add_h1(doc, '10. What may not presently be taught as settled')
add_bullets(doc, [
    'That all religions teach the same metaphysics.',
    'That one living leader has uniquely infallible access to ultimate reality.',
    'That a private revelation binds another adult without ordinary consent and governance.',
    'That Jesus\' divinity/Trinity/atonement mechanics/resurrection metaphysics are established Way doctrine.',
    'That Muhammad\'s final prophethood or Qur\'anic divine authorship are established Way doctrine.',
    'That Buddhist rebirth, nirvana metaphysics, or a final interpretation of non-self are established Way doctrine.',
    'That the Library of Witnesses is a canon.',
    'That a specific afterlife mechanism has been established.',
    'That one marriage/family model is a spiritual loyalty test.',
    'That ordination creates a metaphysically superior caste.',
])

add_h1(doc, '11. Teaching protocol for Houses and Ministers')
add_numbered(doc, [
    'Identify whether a material claim is historical, textual, phenomenological, philosophical, theological, ethical, symbolic, or institutional.',
    'Identify the applicable confidence status where misunderstanding is reasonably likely.',
    'Distinguish primary witness from later interpretation and institutional conclusion.',
    'Represent significant competing interpretations fairly enough that a listener can understand the actual disagreement.',
    'State known limitations and avoid manufactured certainty.',
    'Identify practical and rights implications before a teaching is used to govern behavior.',
    'Record material corrections publicly enough to prevent old errors from persisting by inertia.'
])

add_h1(doc, '12. Review cycle')
add_bullets(doc, [
    'House-level questions and practice feedback are recorded continuously.',
    'The Teaching & Practice Baseline receives formal annual review during the founding period.',
    'Material change requires an evidence packet, competing interpretations, rights-impact review, recorded decision, effective date, and review date.',
    'No teaching becomes entrenched by repetition, popularity, founder preference, donor pressure, or liturgical habit.'
])

baseline_path = OUT/'teaching_and_practice_baseline_v1_0.docx'
doc.save(baseline_path)

# ---------------- Companion Handbook ----------------
doc = set_doc_defaults(Document(), 'Companion Handbook')
title_page(doc, 'Companion Handbook', 'A practical guide to walking the Way, belonging to the Assembly, and living safely in a House of Prayer', 'Founding edition for pilot community use')

add_h1(doc, 'Welcome')
add_para(doc, 'The Way is not primarily something you join. It is a path you may choose to walk. A Companion is simply a person who chooses to walk that path with others while retaining conscience, outside relationships, and the freedom to leave.')
add_callout(doc, 'In one sentence', 'We are Companions of the Way. Together we are the Assembly. We gather in Houses of Prayer for prayer, study, fellowship, discernment, and service - seeking truth without surrendering conscience.')

add_h1(doc, '1. Start here: the whole in plain language')
add_table(doc, ['Term','What it means'], [
    ('the Way','The religious path of seeking truth and living in response to it.'),
    ('Companion','A person who freely walks the Way with others.'),
    ('the Assembly','The wider religious community.'),
    ('House of Prayer','A local community of the Assembly; it need not own a building.'),
    ('Gathering','A specific meeting for prayer, study, service, fellowship, or discernment.'),
    ('Teacher','A person trusted to teach within demonstrated competence.'),
    ('Minister','A person commissioned for defined religious/pastoral functions.'),
    ('Steward','A person entrusted with defined operational or resource responsibilities.'),
    ('Board of Directors','The civil governing board of the nonprofit corporation; not the same thing as the Assembly.')
], widths=[1.7,5.3])

add_h1(doc, '2. You can participate before you belong')
add_para(doc, 'A person may attend public gatherings, study, pray, ask questions, join service projects, eat with the community, disagree, or remain a long-term participant without becoming a Companion. No six-month visitor, lifelong guest, or curious neighbor should be pressured into religious identification.')
add_bullets(doc, [
    'You do not need to donate to attend.',
    'You do not need to disclose private information to prove sincerity.',
    'You do not need to stop reading critics or participating in outside relationships.',
    'You do not need to adopt every provisional teaching.',
    'You may leave at any time.'
])

add_h1(doc, '3. Becoming a Companion')
add_para(doc, 'Ordinarily, become familiar enough with the Way to understand its rights, core practices, teaching-status system, House structure, safeguarding route, and the fact that major metaphysical questions remain open. Then, if you freely choose, make the Companion Commitment publicly, privately, or in writing.')
add_h2(doc, 'The Companion Commitment')
for line in commitment:
    add_para(doc, line)
add_para(doc, 'The Commitment is religious rather than contractual. It does not transfer property, waive legal rights, create a lifetime vow, create statutory corporate membership, or authorize leaders to control your personal decisions.')

add_h1(doc, '4. Your rights')
add_bullets(doc, rights)
add_h2(doc, 'What happens if you disagree?')
add_para(doc, 'Ordinary disagreement is not misconduct. A Companion may criticize a teaching, oppose a leader, decline a rite, reject a Minister\'s advice, consult another religion, or advocate revision. Conduct rules may still address harassment, violence, discrimination, fraud, retaliation, or other real harms, but disagreement itself is not a punishable offense.')

add_h1(doc, '5. Everyday practice')
add_para(doc, 'The Way currently uses a sevenfold practice rather than a long checklist of mandatory observances.')
add_table(doc, ['Practice','Try this in ordinary life'], [
    ('Attention','Create regular moments of stillness, observation, and honest awareness.'),
    ('Prayer','Speak, listen, give thanks, lament, ask, praise, or sit in reverent silence according to conscience.'),
    ('Study','Read primary sources, history, scholarship, and competing interpretations.'),
    ('Discernment','Test motives, claims, choices, and proposed actions rather than treating impulse as revelation.'),
    ('Ethical Practice','Tell the truth, honor consent, restrain harm, keep commitments, repair wrongdoing, and seek justice.'),
    ('Service','Help people and communities without making aid conditional on religious loyalty.'),
    ('Companionship','Share life, hospitality, encouragement, correction, and care with others without ownership.')
], widths=[1.45,5.55])
add_h2(doc, 'A flexible rhythm')
add_bullets(doc, [
    'Daily: some form of attention/prayer and ethical reflection when practicable.',
    'Weekly: some form of communal gathering, rest, study, fellowship, or service.',
    'Periodic: service, Common Table/fellowship, extended study, retreat, remembrance, or discernment.',
    'Seasonal/annual: practices may be piloted, reviewed, and revised rather than imposed permanently from the beginning.'
])

add_h1(doc, '6. What happens at a House of Prayer')
add_para(doc, 'A healthy House should feel recognizably religious without becoming a controlled environment. It should make room for devotion and serious inquiry at the same time.')
add_table(doc, ['Gathering','What to expect'], [
    ('Prayer','Welcome, silence/prayer, readings/reflection, intercession, closing; music or spontaneous prayer may be included.'),
    ('Study','Primary source, context, competing interpretations, discussion, and explicit uncertainty.'),
    ('Service','A real community need addressed voluntarily without conversion requirements.'),
    ('Fellowship / Common Table','Shared food or relationship time without status seating or insider privilege.'),
    ('Discernment','A clearly stated question, conflict disclosure, evidence and arguments, reflection, recorded outcome, and acknowledged uncertainty.')
], widths=[1.6,5.4])
add_h2(doc, 'What should never happen at a Gathering')
add_bullets(doc, [
    'Sleep deprivation, prolonged exhaustion, forced fasting, forced confession, humiliation, intimidation, deceptive recruitment, restricted exit, confiscation of property/devices, or social isolation to produce commitment.',
    'Public spiritual commands over another adult\'s marriage, sexuality, healthcare, finances, residence, employment, education, or family relationships.',
    'Pressure to prove devotion through money, labor, secrecy, loyalty to a leader, or separation from outsiders.'
])

add_h1(doc, '7. How teaching works')
add_para(doc, 'The Way does not treat every religious statement as equally certain. Teachers should distinguish evidence, historical reconstruction, interpretation, theology, hope, symbolism, and institutional decision.')
add_table(doc, ['Label','What it means for you'], [
    ('Established','The Assembly considers the defined claim strongly supported, though still revisable.'),
    ('Well-Supported','Best current conclusion, but important uncertainty remains.'),
    ('Provisional','Useful working conclusion or practice under review.'),
    ('Permitted Interpretation','You may hold it; the Assembly does not require it.'),
    ('Open Question','The Assembly does not presently know enough to settle it.'),
    ('Rejected as Harmful or Unsupported','It may be studied, but cannot be used as an accepted basis for coercion or harm.')
], widths=[1.7,5.3])
add_h2(doc, 'Examples of current open questions')
add_bullets(doc, [
    'The precise metaphysical nature of God/ultimate reality.',
    'Ontological Christology, Trinity, atonement mechanisms, and resurrection metaphysics.',
    'The divine authorship of the Qur\'an and finality of Muhammad\'s prophethood.',
    'Buddhist rebirth, nirvana metaphysics, and final implications of non-self.',
    'Immortal soul, resurrection, rebirth, heaven/hell, annihilation, universal reconciliation, and other afterlife models.',
    'Final marriage theology, ordination theology, comprehensive sacred calendar, and formal doctrine of dual religious belonging.'
])

add_h1(doc, '8. Family, intimate life, and personal decisions')
add_para(doc, 'The founding baseline is intentionally stronger on consent and safety than on unresolved metaphysical claims about family form.')
add_bullets(doc, [
    'Meaningful consent, capacity, non-exploitation, honesty, safeguarding, and responsibility to dependents are required.',
    'Spiritual status, prophecy, money, employment, housing, immigration status, dependency, or threatened exclusion may not be used to obtain sexual or marital compliance.',
    'No person gains spiritual rank by marrying, remaining single, having children, remaining childless, or conforming to a leader\'s preferred family form.',
    'Pastoral care may help you examine values and safety; it may not compel marriage, divorce, reconciliation, reproduction, abstinence from lawful adult relationships, or disclosure of intimate facts merely to prove loyalty.'
])

add_h1(doc, '9. Suffering, healing, and death')
add_bullets(doc, [
    'You should not be told that suffering proves divine punishment, karmic desert, weak faith, or spiritual inferiority.',
    'Prayer for healing is allowed; refusing or seeking medical/mental-health treatment is never a test of faith.',
    'Funerals and memorials should honor the dead, support mourners, and distinguish hope from certainty.',
    'No Minister may claim certain knowledge of another person\'s final destiny merely by office.'
])

add_h1(doc, '10. Ministers, Teachers, Stewards, and Elders')
add_para(doc, 'Leadership in the Way is a function, not a superior species of person. You may ask what authority a person actually has. Their written role, commission, policy, or corporate delegation should answer the question.')
add_bullets(doc, [
    'A Teacher can be corrected.',
    'A Minister can be investigated, suspended, or removed.',
    'A Steward does not own what they administer.',
    'An Elder, if the role is adopted, does not automatically control money, discipline, doctrine, or revelation.',
    'The founder is subject to the same complaint, conflict, compensation, safeguarding, and investigation systems as other leaders.'
])

add_h1(doc, '11. Money, work, housing, and charitable help')
add_bullets(doc, [
    'Giving is voluntary. No donation purchases spiritual standing, office, healing, revelation, protection, or special access.',
    'Employment is employment, not proof of spiritual loyalty. Wages, benefits, supervision, and discipline should follow ordinary lawful processes.',
    'Housing, if the Way later provides it, may not be used to trap people in religious participation or silence criticism.',
    'Charitable assistance is based on charitable purpose and need, not religious conformity, donation history, silence, or membership.'
])

add_h1(doc, '12. Safeguarding, complaints, and retaliation')
add_para(doc, 'If something is unsafe, abusive, coercive, fraudulent, or retaliatory, spiritual reconciliation does not replace ordinary reporting and protection. Child abuse/neglect and other mandatory-reporting matters must be handled under applicable law before internal spiritual processes.')
add_bullets(doc, [
    'Use the normal House contact for ordinary concerns when appropriate.',
    'Use the independent integrity/safeguarding channel when leadership is implicated or independence is needed.',
    'Anonymous reports may be accepted.',
    'An allegation that cannot be substantiated is not automatically a bad-faith allegation.',
    'Retaliation for reporting, witnessing, criticizing, declining participation, or leaving is prohibited.'
])

add_h1(doc, '13. Leaving - and returning')
add_para(doc, 'A clear statement that you no longer identify as a Companion is enough. No permission, exit interview, confession, payment, doctrinal recantation, nondisparagement promise, or spiritual release is required.')
add_bullets(doc, [
    'Friends and relatives are not instructed to shun you.',
    'Confidential pastoral information remains protected after departure.',
    'You may criticize the Way publicly.',
    'Ordinary property, wages, benefits, deposits, records, and contractual rights are handled under ordinary law, not religious standing.',
    'You may later return if you freely wish; return cannot be conditioned on denying legitimate criticism or past experience.'
])

add_h1(doc, '14. A practical first 90 days')
add_table(doc, ['Period','Suggested approach'], [
    ('Weeks 1-4','Attend different gathering types; read the Constitution and Teaching & Practice Baseline; ask difficult questions; keep outside relationships normal.'),
    ('Weeks 5-8','Try a personal rhythm of attention, prayer, study, service, and companionship; participate in a service activity; learn the complaint/safeguarding route.'),
    ('Weeks 9-12','Decide whether the Companion identity honestly describes your relationship to the Way. If yes, make the Commitment in a manner meaningful to you. If no, continue participating or leave without penalty.')
], widths=[1.3,5.7])

add_h1(doc, '15. Glossary')
add_table(doc, ['Term','Short definition'], [
    ('Encounter','A reported experience interpreted as contact with or awareness of transcendent reality.'),
    ('Testimony','An account of an encounter, conviction, event, or practice.'),
    ('Interpretation','The human act of assigning meaning to testimony, text, event, or experience.'),
    ('Teaching','A proposition or practice currently taught by the Assembly.'),
    ('Doctrine','A formally adopted teaching with defined status and implications.'),
    ('Discernment','Testing claims, motives, choices, and interpretations individually and communally.'),
    ('Commissioning','Formal authorization for a defined ministry or service function.'),
    ('Library of Witnesses','The open research corpus studied by the Way; inclusion does not mean canonization or endorsement.'),
    ('Rule of Life','A voluntary, adaptable pattern of practice; never a coercive surveillance code.'),
    ('Former Companion','A person who no longer identifies as a Companion and retains ordinary dignity, relationships, and rights.')
], widths=[1.8,5.2])

add_h1(doc, '16. The maturity test')
add_para(doc, 'A healthy Way is not measured by how difficult it is to leave or how dependent people become on its leaders. It is measured by whether people grow in honesty, compassion, discernment, responsibility, freedom from manipulation, service, and capacity to seek truth - and whether the community can survive correction.')
add_callout(doc, 'Remember', 'You are a Companion, not property of the Assembly. A Minister can serve you without owning your conscience. A House can become home without becoming a cage.', fill='E8F5E9', color=GREEN)

handbook_path = OUT/'companion_handbook_v1_0.docx'
doc.save(handbook_path)

# ---------------- Source artifacts ----------------
constitution_md = '''# Constitution of the Way v1.0\n\nCanonical editable source is represented by the DOCX and the structured JSON/crosswalk in this release. This Markdown file is a release index rather than a line-for-line export.\n\nSee `constitution_of_the_way_v1_0.docx`.\n'''
(SRC/'constitution_of_the_way_v1_0.md').write_text(constitution_md)
(SRC/'teaching_and_practice_baseline_v1_0.md').write_text('# Teaching & Practice Baseline v1.0\n\nSee compiled DOCX/PDF in release root.\n')
(SRC/'companion_handbook_v1_0.md').write_text('# Companion Handbook v1.0\n\nSee compiled DOCX/PDF in release root.\n')

crosswalk = [
    ['C-01','Nature/purpose','v0.5-v0.7','Constitution Articles I-III'],
    ['C-02','Conscience/exit/anti-coercion','Integrity v0.2-v0.4','Constitution Articles II, IV, XVII'],
    ['C-03','Companion pathway','v0.6','Constitution Article V; Companion Handbook'],
    ['C-04','Assembly/Houses','v0.6','Constitution Articles VI-VII'],
    ['C-05','Ministry/commissioning','v0.6-v0.9','Constitution Article VIII; Teaching Baseline'],
    ['C-06','Epistemic method','v0.5, v0.7-v0.8','Constitution Articles IX-X'],
    ['C-07','Practice/formation','v0.7','Constitution Article XI; Teaching Baseline; Handbook'],
    ['C-08','Foundational theology','v0.8','Teaching Baseline Section 3'],
    ['C-09','Ethics/eschatology','v0.9','Constitution Articles XII-XIV; Teaching Baseline Sections 4, 7, 9'],
    ['C-10','Founder/succession','v0.2-v0.6','Constitution Article XVI'],
]
with open(SRC/'constitutional_crosswalk_v1_0.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['ID','Domain','Source lineage','v1.0 location']); w.writerows(crosswalk)

teaching_register = core_teachings + [
    ('TB-016','Suffering','Established','Do not presume suffering proves punishment, desert, weak faith, or inferiority.'),
    ('TB-017','Theodicy','Open Question','No single theodicy adopted.'),
    ('TB-018','Spiritual end','Provisional','Healing/liberation/reconciliation/awakening/wholeness/right relationship as working comparative terms.'),
    ('TB-019','Exclusive salvation','Open Question','No claim that only Companions reach the ultimate good.'),
    ('TB-020','Intimate ethics','Established','Capacity, consent, non-exploitation, honesty, safeguarding, responsibility.'),
    ('TB-021','Marriage theology','Open Question','No final doctrine yet.'),
    ('TB-022','Violence','Well-Supported','Strong presumption toward nonviolence/peacemaking/protection; absolute pacifism open.'),
    ('TB-023','Religious violence','Established','No vigilantism/holy war/coercive enforcement/revelatory exemption from law.'),
    ('TB-024','Commissioning','Established','Defined, reviewable functional authorization; no sacred caste.'),
    ('TB-025','Ordination','Reserved','Not adopted unless future theology requires more than commissioning.'),
    ('TB-026','Sacred time','Provisional','Rhythm before comprehensive calendar.'),
    ('TB-027','Interreligious posture','Established','Disciplined hospitality; accurate representation; no compelled conversion.'),
    ('TB-028','Dual belonging','Open Question','Case-specific honesty; no blanket rule yet.'),
]
with open(SRC/'teaching_status_register_v1_0.csv','w',newline='') as f:
    w=csv.writer(f); w.writerow(['ID','Domain','Status','Statement']); w.writerows(teaching_register)

structure = {
  'version':'1.0',
  'date':'2026-09-20',
  'identity':{'path':'the Way','person':'Companion','community':'the Assembly','local':'House of Prayer'},
  'entrenched_principles':['conscience','voluntary participation','safe exit','non-retaliation','safeguarding','no infallible living authority','no founder supremacy','financial integrity','due process','religious/civil authority distinction'],
  'epistemic_chain':['Encounter','Testimony','Interpretation','Teaching','Doctrine'],
  'confidence_states':['Established','Well-Supported','Provisional','Permitted Interpretation','Open Question','Rejected as Harmful or Unsupported'],
  'sevenfold_practice':['Attention','Prayer','Study','Discernment','Ethical Practice','Service','Companionship'],
  'documents':['Constitution of the Way','Teaching & Practice Baseline','Companion Handbook']
}
(SRC/'way_foundational_synthesis_v1_0.json').write_text(json.dumps(structure,indent=2))

(SRC/'decision_register_v1_0.md').write_text('''# v1.0 Decision Register\n\n## Adopted in founding synthesis\n- The Constitution of the Way is a religious constitution distinct from civil charter/bylaws.\n- Entrenched rights and anti-coercion principles sit above ordinary teaching/practice.\n- Companion is religious identity, not statutory corporate membership.\n- Commissioning remains the present ministry authorization model; ordination remains reserved.\n- Library of Witnesses remains open; no Way canon is declared.\n- Sevenfold practice is the founding practice framework.\n- Teaching status grammar is constitutionalized for the founding period.\n- The Assembly may teach ethical safeguards more firmly than unresolved metaphysical claims.\n- Founder status remains historical only; succession is system transmission, not appointment of an heir.\n\n## Deliberately unresolved\n- Final legal/public corporate name.\n- Comprehensive metaphysics of God/ultimate reality.\n- Christology, Qur'anic divine authorship/finality, Buddhist rebirth/nirvana metaphysics.\n- Final afterlife doctrine.\n- Comprehensive marriage/family theology.\n- Ordination theology.\n- Comprehensive sacred calendar.\n- Formal doctrine of dual religious belonging.\n''')

readme = '''# The Way v1.0 Founding Synthesis\n\nThis release reconciles v0.5-v0.9 into the first coherent public-facing religious baseline.\n\n## Core publications\n1. Constitution of the Way v1.0\n2. Teaching & Practice Baseline v1.0\n3. Companion Handbook v1.0\n\n## Status\nFounding synthesis for pilot use, community review, legal/corporate reconciliation, and future amendment. It is not a claim that theology is complete.\n\n## Critical distinction\nThe Constitution of the Way is a religious constitution. It does not replace the Tennessee nonprofit corporation's charter/bylaws or civil-law fiduciary responsibilities.\n'''
(OUT/'README_RELEASE.md').write_text(readme)
print(constitution_path)
print(baseline_path)
print(handbook_path)
